# pptx 근사 미리보기 — LibreOffice Impress 가 없는 환경에서 배치·겹침·넘침을 눈으로 확인하기 위한 렌더러
# (python-pptx 로 도형 좌표·텍스트·이미지를 읽어 PIL 로 그린다. 글꼴 폭은 근사)
import sys, io, textwrap
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageDraw, ImageFont

SRC = sys.argv[1]; OUT = sys.argv[2]; DPI = float(sys.argv[3]) if len(sys.argv) > 3 else 72
FONT = '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc'
prs = Presentation(SRC)
SW, SH = prs.slide_width, prs.slide_height
px = lambda emu: int(Emu(emu).inches * DPI)
W, H = px(SW), px(SH)
_fonts = {}
def font(pt, bold=False):
    k = (round(pt * DPI / 72), bold)
    if k not in _fonts: _fonts[k] = ImageFont.truetype(FONT, max(6, k[0]), index=0)
    return _fonts[k]

def color_of(fill):
    try:
        if fill.type == 1 and fill.fore_color and fill.fore_color.rgb: return '#' + str(fill.fore_color.rgb)
    except Exception: pass
    return None

def draw_text(d, shape, x, y, w, h):
    tf = shape.text_frame
    ins = px(91440 * 0.05) if shape.text_frame.margin_left == 0 else px(shape.text_frame.margin_left)
    cy = y + px(shape.text_frame.margin_top)
    lines_total = []
    for para in tf.paragraphs:
        runs = para.runs
        if not runs: lines_total.append(('', 12, False, '#000000', 0)); continue
        size = runs[0].font.size.pt if runs[0].font.size else 12
        bold = bool(runs[0].font.bold)
        col = '#000000'
        try:
            if runs[0].font.color and runs[0].font.color.rgb: col = '#' + str(runs[0].font.color.rgb)
        except Exception: pass
        text = ''.join(r.text for r in runs)
        bullet = para._p.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}buChar') is not None
        f = font(size, bold)
        avail = w - 2 * ins - (px(0.2 * 914400) if bullet else 0)
        # 줄바꿈: 픽셀 폭 기준
        words, line, out = list(text), '', []
        for ch in words:
            if d.textlength(line + ch, font=f) > avail and line: out.append(line); line = ch
            else: line += ch
        out.append(line)
        for i, ln in enumerate(out): lines_total.append((('• ' if bullet and i == 0 else ('  ' if bullet else '')) + ln, size, bold, col, para.space_after.pt if para.space_after else 0))
    total_h = sum(px(s * 914400 / 72 * 1.25) + px(sa * 914400 / 72) for _, s, _, _, sa in lines_total)
    align = tf.paragraphs[0].alignment if tf.paragraphs else None
    anchor = tf.vertical_anchor
    if anchor is not None and int(anchor) == 3: cy = y + (h - total_h) // 2         # middle
    for text, size, bold, col, sa in lines_total:
        f = font(size, bold); lw = d.textlength(text, font=f)
        tx = x + ins
        if align is not None and int(align) == 2: tx = x + (w - lw) / 2
        elif align is not None and int(align) == 3: tx = x + w - ins - lw
        d.text((tx, cy), text, font=f, fill=col)
        cy += px(size * 914400 / 72 * 1.25) + px(sa * 914400 / 72)
    return cy - y  # 사용한 높이

overflow = []
for i, slide in enumerate(prs.slides, 1):
    img = Image.new('RGB', (W, H), '#FFFFFF')
    try:
        bg = slide.background.fill
        c = color_of(bg)
        if c: img = Image.new('RGB', (W, H), c)
    except Exception: pass
    d = ImageDraw.Draw(img)
    for sh in slide.shapes:
        x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                im = Image.open(io.BytesIO(sh.image.blob)).convert('RGB').resize((max(1, w), max(1, h)))
                img.paste(im, (x, y))
            except Exception as e: d.rectangle([x, y, x + w, y + h], outline='#FF00FF')
            continue
        if sh.shape_type == MSO_SHAPE_TYPE.TABLE:
            d.rectangle([x, y, x + w, y + h], outline='#999999')
            tbl = sh.table; rh = h / max(1, len(tbl.rows)); cy = y
            for r in tbl.rows:
                cx = x
                for ci, cell in enumerate(r.cells):
                    cw = px(tbl.columns[ci].width)
                    d.rectangle([cx, cy, cx + cw, cy + rh], outline='#BBBBBB')
                    t = cell.text[:40]; d.text((cx + 3, cy + 2), t, font=font(9), fill='#000000')
                    cx += cw
                cy += rh
            continue
        fillc = color_of(sh.fill) if hasattr(sh, 'fill') else None
        if fillc: d.rectangle([x, y, x + w, y + h], fill=fillc)
        if sh.has_text_frame and sh.text_frame.text.strip():
            used = draw_text(d, sh, x, y, w, h)
            if used > h + px(0.15 * 914400): overflow.append((i, sh.text_frame.text[:30], round(used / DPI, 2), round(h / DPI, 2)))
    img.save(f'{OUT}/s{i:02d}.png')
print('rendered', len(prs.slides))
for o in overflow: print('OVERFLOW?', o)
