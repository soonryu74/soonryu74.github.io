"""IEEE AgeTech 2026 dataset report (docx). Times New Roman 11 pt, 1.15 spacing, single column."""
import sys, json, pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
R = sys.argv[1] if len(sys.argv) > 1 else 'khalo'
AUTHOR, EMAIL = sys.argv[2], sys.argv[3]
qc = json.load(open(f'{R}/qc/qc_summary.json'))
t1 = pd.read_csv(f'{R}/benchmarks/results/t1_forecasting.csv'); t2 = pd.read_csv(f'{R}/benchmarks/results/t2_hle_regression.csv')
t3 = pd.read_csv(f'{R}/benchmarks/results/t3_equity_audit.csv', index_col=0)
ind = pd.read_csv(f'{R}/data/indicators.csv')
doc = Document()
st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.line_spacing = 1.15; st.paragraph_format.space_after = Pt(4)
for s in doc.sections: s.top_margin = s.bottom_margin = Cm(2.2); s.left_margin = s.right_margin = Cm(2.3)
for lvl in (1, 2):
    h = doc.styles[f'Heading {lvl}']; h.font.name = 'Times New Roman'; h.font.color.rgb = RGBColor(0, 0, 0)
    h.font.size = Pt(13 if lvl == 1 else 11.5); h.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    h.paragraph_format.space_before = Pt(10 if lvl == 1 else 6); h.paragraph_format.space_after = Pt(3)
P = lambda t, **k: para(t, **k)
def para(t, bold=False, italic=False, size=None, align=None, after=None):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p
def rich(parts):
    p = doc.add_paragraph()
    for t, b in parts: r = p.add_run(t); r.bold = b
    return p
def bullet(t, b=None):
    p = doc.add_paragraph(style='List Bullet')
    if b: r = p.add_run(b); r.bold = True
    p.add_run(t); p.paragraph_format.space_after = Pt(1); return p
def table(rows, header, widths=None, size=9):
    tb = doc.add_table(rows=1, cols=len(header)); tb.style = 'Table Grid'
    for i, h in enumerate(header):
        c = tb.rows[0].cells[i]; c.text = ''; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size)
    for row in rows:
        cs = tb.add_row().cells
        for i, v in enumerate(row):
            cs[i].text = ''; r = cs[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
    if widths:
        from docx.oxml import OxmlElement
        tb.autofit = False
        tblPr = tb._tbl.tblPr; lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
        for i, gc in enumerate(tb._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'), str(int(widths[i] * 567)))
        for row in tb.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tb
def cap(t): para(t, italic=True, size=9.5, after=6)

# ---------------- Title page ----------------
for _ in range(5): doc.add_paragraph()
para('K-HALO: Korea Healthy Ageing and Longevity District Panel, 2008–2025', bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
para('An AI-ready longitudinal dataset of small-area health behaviours, chronic disease, preventive care and healthy life expectancy in a super-aged society', italic=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=30)
para(AUTHOR, bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para('Independent researcher, Seoul, Republic of Korea', size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para(f'Contact: {EMAIL}', size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=30)
para('IEEE AgeTech Global Dataset Competition 2026 — Theme 1: Health & Well-Being (chronic disease management; preventive healthcare and wellness)', size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
para('Dataset version 1.0 · Licence CC BY 4.0 · Formats CSV, SQLite, Croissant JSON-LD', size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ---------------- 1 Abstract ----------------
doc.add_heading('1. Abstract', 1)
best1 = t1.groupby('model').skill_vs_persistence.mean()
P('K-HALO is an openly licensed, English-documented panel that follows every Community Health Survey (CHS) area of South Korea — 281 health-centre areas, of which 245 form a non-overlapping national benchmark set — over 18 years (2008–2025). '
  'It contains 154,691 crude and age-standardised values of 41 indicators covering smoking, alcohol, physical activity, obesity, mental health, oral health (including chewing difficulty among adults aged 65+), hypertension and diabetes (diagnosis, treatment and complication screening), preventive care, unmet medical need and self-rated health. '
  'Each area is linked to life expectancy and healthy life expectancy (HLE; 2017–2023), an area deprivation index and its older-adult population. '
  f'All values are aggregate official statistics, so the dataset contains no personal information. Within published years, {qc["completeness_benchmark_units_pct"]}% of benchmark cells are complete. '
  f'We define three reproducible benchmark tasks — next-year forecasting of ten ageing-related indicators, HLE regression with leave-one-province-out validation, and an equity audit across deprivation quintiles — and report baselines: '
  f'a ridge model improves on persistence by {best1.max()*100:.0f}% on average, explains {t2.R2.max():.2f} of between-district HLE variance, and forecasting error in the most deprived fifth of areas is {(t3.loc["gradient boosting (lags of all 41 indicators)","Q5_over_Q1"]-1)*100:.0f}% higher than in the least deprived fifth. '
  'K-HALO gives AgeTech researchers a ready testbed for small-area forecasting, healthy-longevity modelling and fairness evaluation in the world\'s fastest-ageing population.')

# ---------------- 2 Motivation ----------------
doc.add_heading('2. Motivation', 1)
P('South Korea crossed the super-aged threshold (20% of residents aged 65+) in December 2024 and is projected to have one of the oldest populations in the world within two decades. '
  'Healthy longevity, however, is unevenly distributed: in K-HALO, district HLE ranges from 63.9 to 78.6 years, and in 2023 lower HLE was associated with higher area deprivation (Pearson r = −0.32 across 227 districts; Figure 2). '
  'Decisions about where to deploy AgeTech — remote monitoring, preventive programmes, chronic-disease management, community care — are made for places, not only for individuals.')
P('Most open AgeTech datasets are sensor- or clinic-based samples of tens to hundreds of participants, typically from Western countries. They are essential for device-level AI but cannot tell planners which communities are ageing in poor health, how fast risk factors are changing, or whether an AI model is equally accurate for deprived and affluent areas. '
  'Korea publishes rich small-area statistics, but they are scattered across hundreds of Korean-language tables, change format over time, and are rarely used in machine-learning research.')
P('K-HALO closes this gap by (i) harmonising 18 years of small-area statistics into a stable panel with consistent identifiers, (ii) linking them to healthy life expectancy, deprivation and older-adult population, (iii) documenting every field in English with a datasheet and Croissant metadata, and (iv) defining benchmark tasks with fixed splits, baselines and a built-in equity audit.')

# ---------------- 3 Data Collection ----------------
doc.add_heading('3. Data Collection', 1)
P('K-HALO is a secondary compilation; no new data were collected from people. All inputs are official aggregate statistics distributed through KOSIS, the national statistical portal, and were retrieved through the KOSIS Open API in September 2026.')
table([
 ['Community Health Survey (KDCA)', 'Annual, 2008–2025; ≈230,000 adults/year, ≈900 per health-centre area; face-to-face interviews; area-level crude and age-standardised rates', '41 indicators (chs_panel_*.csv)'],
 ['Cause-of-death statistics (Statistics Korea)', 'Deaths by district, sex and 5-year age group', 'Life tables → life expectancy (hle.csv)'],
 ['Resident registration population (MOIS / Statistics Korea)', 'Mid-year population by district and 5-year age', 'Life tables; older-adult counts (population_age.csv)'],
 ['Population and Housing Census (Statistics Korea)', '2020 district aggregates of education, occupation, household and housing', 'Deprivation index (deprivation.csv)'],
], ['Source', 'Collection', 'Used for'], widths=[4.5, 7.5, 4.5])
cap('Table 1. Data sources. Devices, sensors and recruitment are not applicable: the CHS uses probability sampling of households by trained interviewers, and K-HALO uses only its published area estimates.')
doc.add_heading('3.1 Collection methodology', 2)
bullet(' An inventory of 206 KOSIS tables was screened automatically for coverage (years, areas, missingness). Of 87 CHS tables, 55 passed and 32 were flagged (mostly discontinued series); none failed. The 41 indicators still published by KDCA were retained.', 'Selection.')
bullet(' Areas follow the CHS sampling units (health-centre areas). Administrative changes (city mergers, new health centres) are recorded with status, first and last year; Statistics Korea 2018 district codes are provided for linkage to other data.', 'Harmonisation.')
bullet(' District abridged life tables were built from 3-year pooled deaths and population with Chiang\'s method (age groups 0–4 … 85+; areas with <150 deaths in the window are not estimated). HLE uses the Sullivan method: the national age-specific unhealthy share is back-calculated from Statistics Korea\'s official subjective-health HLE, and each area\'s curve is shifted by the odds ratio of its 3-year mean self-rated health. Re-computed national HLE matches the official value to 0.1 year in all seven published years (2012–2024).', 'Life expectancy and HLE.')
bullet(' A reproduction of the Korean area deprivation index (Kim et al., 2013): z-scores of seven 2020 census variables (low education, aged 65+, low-skill occupation, one-person households, non-apartment housing, female-headed households, divorced/widowed) are summed; car ownership and substandard housing were unavailable.', 'Deprivation.')

# ---------------- 4 Characteristics ----------------
doc.add_heading('4. Dataset Characteristics', 1)
table([
 ['Spatial units', '281 health-centre areas (235 districts, 46 sub-districts); benchmark set 245 (74 metropolitan, 84 urban, 76 rural, 11 mixed)'],
 ['Time span', '2008–2025 (18 survey years); HLE 2017–2023; deprivation 2020; population 2024'],
 ['Indicators', '41 indicators in 9 domains; crude and age-standardised rate (%), direction flag (higher is worse)'],
 ['Samples', '154,691 unit-year-indicator values; 4,878 unit-years in the wide ML file; 1,763 HLE unit-years'],
 ['Participants', 'Not applicable (aggregate data); underlying survey ≈230,000 adults per year'],
 ['Labels / targets', 'Task-defined: next-year rate (T1), HLE (T2), deprivation quintile for stratified evaluation (T3)'],
 ['File formats', 'CSV (UTF-8, long and wide), SQLite, Croissant 1.0 JSON-LD, data dictionary CSV, Markdown datasheet'],
], ['Property', 'Value'], widths=[3.6, 12.9])
cap('Table 2. Dataset summary.')
dom = ind.groupby('domain_en').agg(n=('indicator_id', 'size'), ex=('name_en', lambda s: '; '.join(list(s)[:3]))).reset_index()
table([[r.domain_en, r.n, r.ex] for r in dom.itertuples()], ['Domain', 'n', 'Example indicators'], widths=[3.4, 1.0, 12.1], size=8.5)
cap('Table 3. Indicator domains (full list with Korean names, denominators and years in indicators.csv and Appendix A).')
P('Tables are linked by unit_id. The wide file (one row per unit-year, one column per indicator) is directly usable for supervised learning; the long file supports time-series and panel models. Sub-districts carry parent_unit_id so that users can aggregate or join census covariates.')

# ---------------- 5 Data Quality ----------------
doc.add_heading('5. Data Quality', 1)
bullet(f' No values are imputed. Missing cells occur only where KDCA did not survey an indicator in a given year or an area did not yet exist; within published years {qc["completeness_benchmark_units_pct"]}% of benchmark-unit cells are filled (minimum per indicator 99.2%). Figure 1 shows coverage by year.', 'Missing data.')
bullet(f' Range checks found {qc["rate_out_of_0_100"]} rates outside 0–100%. Crude and age-standardised rates agree closely (r = {qc["crude_vs_std_pearson_r"]}, median absolute difference {qc["crude_vs_std_median_abs_diff_pp"]} pp). Year-on-year changes larger than 4 SD of each indicator\'s change distribution affect {qc["yoy_jumps_over_4sd_pct"]}% of {qc["yoy_pairs"]:,} pairs; they are retained and flagged in the QC script because they reflect genuine small-area sampling variation. HLE never exceeds life expectancy (0 violations); LE ranges {qc["le_range"][0]}–{qc["le_range"][1]} and HLE {qc["hle_range"][0]}–{qc["hle_range"][1]} years.', 'Validation.')
bullet(' Retrieval was validated table by table for year and area coverage (206 KOSIS tables); national HLE re-computation reproduces official statistics; all scripts are deterministic and included (tools/, qc/), so the package can be rebuilt and re-validated.', 'Quality assurance.')
doc.add_picture(f'{R}/qc/figures/fig1_coverage.png', width=Cm(14.5))
cap('Figure 1. Share of the 245 benchmark units with a value, by indicator and year. White cells are years in which KDCA did not field the item.')

# ---------------- 6 Ethics ----------------
doc.add_heading('6. Ethical Considerations', 1)
bullet(' K-HALO contains only area-level aggregates already published by government agencies; no individual, household or facility records. Each estimate summarises roughly 900 survey respondents, so re-identification is not possible.', 'Privacy protection.')
bullet(' Not applicable to K-HALO. The CHS is a statutory survey under the Regional Public Health Act, conducted by trained interviewers with respondents\' agreement; only its published area-level estimates are used. No human-subject contact occurred in compiling K-HALO, so no additional ethics review was required.', 'Consent.')
bullet(' Aggregation is the anonymisation step; no small-cell values below the producers\' publication thresholds are introduced. Life-table estimates are suppressed for areas with fewer than 150 deaths in three years.', 'Anonymisation.')
bullet(' Source data are open under the Korea Open Government Licence Type 1 (attribution); K-HALO is released under CC BY 4.0 with full attribution. Derived HLE and deprivation values are labelled as research estimates, not official statistics.', 'Regulatory compliance and licensing.')
bullet(' Survey measures are self-reported; HLE is driven mainly by self-rated health; rural and deprived areas show larger forecast errors (Section 7). The equity audit (T3) is part of the benchmark so that model developers must report subgroup performance. K-HALO must not be used to stigmatise communities; the datasheet lists intended and out-of-scope uses.', 'Bias awareness.')

# ---------------- 7 Applications & benchmarks ----------------
doc.add_heading('7. Potential Applications and Benchmark Baselines', 1)
P('Intended AI and AgeTech use cases include: early warning of rising chronic-disease and risk-factor burden in specific communities; targeting of preventive and remote-monitoring programmes; explaining HLE gaps for healthy-longevity policy; transfer learning and multi-task models for small-area health; and fairness evaluation of public-health AI. Three tasks with fixed splits are provided (benchmarks/).')
targets = {'DT_H_OR_INCONV': 'Chewing difficulty (65+)', 'DT_HYPER_DOCTOR': 'Hypertension prevalence', 'DT_HYPER_DOCTOR_DRUG': 'Hypertension treatment',
           'DT_DIA_DOCTOR': 'Diabetes prevalence', 'DT_DIA_TREAT': 'Diabetes treatment', 'DT_H_MENTAL_DEPRESS': 'Depressive mood',
           'DT_H_EX_WALK': 'Walking practice', 'DT_H_OBE_OBE': 'Obesity', 'DT_HEALTH_RECOG': 'Good self-rated health', 'DT_INFLUENZA': 'Influenza vaccination'}
pv = t1.pivot(index='target', columns='model', values='MAE_pp')
M = ['persistence (last value)', 'mean of last 3 years', 'ridge (lags of all 41 indicators)', 'gradient boosting (lags of all 41 indicators)']
rows = [[targets[t]] + [f'{pv.loc[t, m]:.2f}' for m in M] for t in targets]
rows.append(['Mean'] + [f'{pv[m].mean():.2f}' for m in M])
table(rows, ['T1 target (age-standardised %)', 'Persistence', 'Mean last 3 y', 'Ridge', 'Gradient boosting'], widths=[5.3, 2.6, 2.8, 2.6, 3.2], size=8.5)
cap(f'Table 4. T1 one-step-ahead forecasting, test years 2024–2025 (about {int(t1.n_test.median())} unit-years per target; MAE in percentage points). Models use lags 1–3 of all 41 indicators and area type; trained on target years 2010–2023.')
table([[r.model, r.n, f'{r.MAE_years:.2f}', f'{r.R2:.2f}'] for r in t2.itertuples()], ['T2 model', 'n', 'MAE (years)', 'R²'], widths=[7, 2, 3, 3], size=9)
cap('Table 5. T2 healthy life expectancy (2023) regression with leave-one-province-out validation (17 folds). Self-rated health is excluded from the features because it is an input to the HLE estimate.')
t3r = t3.reset_index()
table([[r[0], r.task.replace('T3 ', ''), f'{r.Q1:.2f}', f'{r.Q3:.2f}', f'{r.Q5:.2f}', f'{r.Q5_over_Q1:.2f}'] for r in t3r.itertuples(index=False)],
      ['Model', 'Evaluated on', 'Q1 (least)', 'Q3', 'Q5 (most)', 'Q5/Q1'], widths=[5.2, 4.2, 1.8, 1.6, 1.8, 1.6], size=8.5)
cap('Table 6. T3 equity audit: mean absolute error by area deprivation quintile.')
P(f'Findings. A ridge model using the full indicator profile beats naive persistence for nine of ten targets, with one tie (mean skill {best1.max()*100:.0f}%), but errors remain several percentage points, leaving clear headroom for spatio-temporal and multi-task models. HLE is moderately predictable from determinants across unseen provinces (R² = {t2.R2.max():.2f}). '
  'Crucially, forecasting error is higher in the most deprived areas for both naive and learned models — precisely where AgeTech interventions are most needed — making K-HALO a natural benchmark for fairness-aware methods.')
doc.add_picture(f'{R}/qc/figures/fig2_hle_deprivation.png', width=Cm(12.5))
cap('Figure 2. District healthy life expectancy vs. area deprivation, 2023; colour shows life expectancy.')

# ---------------- 8 Accessibility ----------------
doc.add_heading('8. Accessibility and Reproducibility', 1)
bullet(' IEEE DataPort (open access). Package: data/ (7 CSV files + SQLite), metadata/data_dictionary.csv, croissant.json, DATASHEET.md, README.md, LICENSE.md, CITATION.cff, benchmarks/, qc/, tools/.', 'Access.')
bullet(' Standard open formats (CSV UTF-8, SQLite); bilingual names (English/Korean); Statistics Korea codes for spatial joins; MLCommons Croissant 1.0 metadata with Responsible-AI fields for automatic loading in ML frameworks.', 'Interoperability.')
bullet(' python benchmarks/run_baselines.py data reproduces Tables 4–6 (pandas, scikit-learn; fixed seed). python qc/run_qc.py data reproduces Section 5 and the figures. tools/build_khalo.py rebuilds the data from the source statistics.', 'Replication.')
bullet(' Annual updates after each CHS release, with versioned changelog.', 'Maintenance.')

# ---------------- 9 References ----------------
doc.add_heading('9. References', 1)
for r in ['Chiang CL. The Life Table and Its Applications. Malabar, FL: Krieger; 1984.',
          'Gebru T, Morgenstern J, Vecchione B, et al. Datasheets for datasets. Communications of the ACM. 2021;64(12):86–92.',
          'Kang YW, Ko YS, Kim YJ, et al. Korea Community Health Survey data profiles. Osong Public Health and Research Perspectives. 2015;6(3):211–217.',
          'Kim DJ, et al. Korean area deprivation index for health equity monitoring (in Korean). Korea Institute for Health and Social Affairs; 2013.',
          'Korea Disease Control and Prevention Agency. Community Health Survey statistics, 2008–2025. KOSIS, https://kosis.kr (accessed Sep 2026).',
          'MLCommons. Croissant format specification, version 1.0. 2024. https://mlcommons.org/croissant',
          'Statistics Korea. Life tables, cause-of-death statistics and resident registration population. KOSIS, https://kosis.kr (accessed Sep 2026).',
          'Sullivan DF. A single index of mortality and morbidity. HSMHA Health Reports. 1971;86(4):347–354.',
          'United Nations, Department of Economic and Social Affairs. World Population Prospects 2024. New York: UN; 2024.']:
    p = doc.add_paragraph(r); p.paragraph_format.space_after = Pt(2); p.runs[0].font.size = Pt(10)

# ---------------- Appendix ----------------
doc.add_page_break()
doc.add_heading('Appendix A. Indicator list', 1)
table([[r.indicator_id, r.name_en, r.population, f'{r.first_year}–{r.last_year}', {1: 'worse', 0: 'better'}.get(r.higher_is_worse, '–') if pd.notna(r.higher_is_worse) else '–'] for r in ind.itertuples()],
      ['ID (KOSIS table)', 'Indicator', 'Population', 'Years', 'Higher is'], widths=[4.4, 6.6, 3.0, 1.8, 1.4], size=7.5)
doc.save(f'{R}/report/K-HALO_dataset_report.docx'); print('docx ok')
