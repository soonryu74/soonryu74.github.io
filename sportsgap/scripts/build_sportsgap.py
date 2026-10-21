#!/usr/bin/env python3
"""SportsGap AI 데이터 파이프라인.

원자료(sportsgap/raw/) → 화면용 JSON(sportsgap/data/).
저장소 루트에서 실행:  python3 sportsgap/scripts/build_sportsgap.py

입력
  raw/mcst_public_sports_2024.xlsx       문화체육관광부 「2025년 전국 공공체육시설 현황(2024.12.31 기준)」
  raw/jumin_gg_age5_202609.csv           행정안전부 주민등록 인구통계 연령별(5세) 2026년 9월, 전체 시군구
  raw/kostat_2018_municipalities_topo_simple.json  통계청 센서스 행정구역경계(2018) — southkorea-maps 배포본
  raw/bus_stops_gg_20251031.csv          국토교통부 전국 버스정류장 위치정보(2025-10-31)의 경기도 추출본

출력
  data/regions.json       경기도 31개 시군 지표(원값만, 점수 계산은 js/score.js)
  data/gg-sigun.geojson   경기도 시군 경계(구 단위 폴리곤에 sigun 속성 부여)
  data/facilities_gg.csv  경기도 공공체육시설 목록(시군·유형·시설명) — 재현·검증용

점수·가중치 계산은 브라우저(js/score.js)에서 하며, 이 스크립트는 원값만 만든다.
"""
import csv
import json
import math
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(ROOT, "raw")
OUT = os.path.join(ROOT, "data")

SIGUN = [
    "수원시", "용인시", "고양시", "성남시", "화성시", "부천시", "남양주시", "안산시", "평택시", "안양시",
    "시흥시", "김포시", "파주시", "의정부시", "광주시", "광명시", "하남시", "군포시", "오산시", "양주시",
    "이천시", "구리시", "안성시", "의왕시", "포천시", "양평군", "여주시", "동두천시", "과천시", "가평군", "연천군",
]
# 설치주체별 '군' 명칭 표기 차이 흡수
ALIAS = {"광주": "광주시", "여주군": "여주시", "양주군": "양주시", "포천군": "포천시", "화성군": "화성시"}

# 시설 유형 시트 → 화면 표시 이름 (공공체육시설 분류 23항목)
SHEETS = {
    "1.육상경기장": "육상경기장",
    "2.축구장": "축구장",
    "3.하키장": "하키장",
    "4.야구장": "야구장",
    "5.싸이클경기장": "싸이클경기장",
    "6.테니스장": "테니스장",
    "7.씨름장": "씨름장",
    "9.1 구기체육관": "구기체육관",
    "9.2 투기체육관": "투기체육관",
    "9.3 생활체육관": "생활체육관",
    "10. 게이트볼장": "전천후게이트볼장",
    "11. 수영장": "수영장",
    "12. 롤러스케이트장": "롤러스케이트장",
    "13. 사격장": "사격장",
    "14. 국궁장": "국궁장",
    "15. 양궁장": "양궁장",
    "16. 승마장": "승마장",
    "17. 골프연습장": "골프연습장",
    "18. 조정카누장": "조정카누장",
    "19. 요트장": "요트장",
    "20. 빙상장": "빙상장",
    "21.1 스키점프경기장": "설상경기장",
    "21.2 바이애슬론경기장": "설상경기장",
    "21.3 크로스컨트리경기장": "설상경기장",
    "21.4 봅슬레이,루지,스켈레톤경기장": "설상경기장",
    "22. 파크골프장": "파크골프장",
    "기타 체육시설(풋살장)": "풋살장",
    "기타 체육시설(그외)": "기타체육시설",
}
# 생활권에서 일상적으로 쓰는 '생활밀착형' 시설 (시나리오 모드의 시설 유형 필터)
LIFE_TYPES = {"생활체육관", "구기체육관", "수영장", "축구장", "풋살장", "테니스장", "전천후게이트볼장", "파크골프장", "기타체육시설", "투기체육관", "골프연습장"}
ELDER_TYPES = {"전천후게이트볼장", "파크골프장", "생활체육관", "수영장"}
CHILD_TYPES = {"수영장", "축구장", "풋살장", "생활체육관", "구기체육관", "빙상장", "롤러스케이트장"}


def norm(s):
    if s is None:
        return ""
    return re.sub(r"\s+", "", str(s))


def match_sigun(text):
    """'고양시\n덕양구', '경기도 수원시 장안구 ...' 등에서 시군명을 찾는다."""
    t = norm(text)
    if not t:
        return None
    for name in SIGUN:
        if name in t:
            return name
    for a, name in ALIAS.items():
        if a in t:
            return name
    return None


# ---------------------------------------------------------------- 시설
def load_facilities(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows_out = []
    unmatched = []
    for sheet, ftype in SHEETS.items():
        ws = wb[sheet]
        header = None
        cur_sido = None
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=12, values_only=True)):
            cells = [norm(c) for c in row]
            if header is None:
                if "시도" in cells and "시설명" in cells:
                    header = {k: cells.index(k) for k in cells if k}
                    c_sido = header["시도"]
                    c_sgg = header.get("시군구", header.get("시ㆍ군ㆍ구"))
                    c_addr = header["주소"]
                    c_name = header["시설명"]
                continue
            sido, sgg, addr, name = cells[c_sido], cells[c_sgg], cells[c_addr], cells[c_name]
            # 블록 시작(소계/합계 행)에서 시도 갱신
            if sgg in ("소계", "계", "합계") or (sido and sgg in ("소계", "계")):
                cur_sido = sido or cur_sido
                continue
            if sido and sido not in ("전국",):
                cur_sido = sido
            if not name or name.startswith("=") or sgg in ("소계", "계", "합계"):
                continue
            if cur_sido != "경기":
                continue
            sg = match_sigun(sgg) or match_sigun(addr)
            if not sg:
                unmatched.append((sheet, sgg, addr, name))
                continue
            rows_out.append({"sigun": sg, "type": ftype, "name": str(row[c_name]).strip(), "address": str(row[c_addr] or "").strip(), "sheet": sheet})
    # 간이운동장(09경기 시트, 시군별 계)
    simple = {}
    ws = wb["09경기"]
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, max_col=2, values_only=True):
        sg = match_sigun(row[0])
        if sg and isinstance(row[1], (int, float)):
            simple[sg] = int(row[1])
    return rows_out, simple, unmatched


# ---------------------------------------------------------------- 인구
def load_population(path):
    raw = open(path, "rb").read().decode("cp949")
    rdr = csv.reader(raw.splitlines())
    header = next(rdr)
    month = re.match(r"(\d{4}년\d{2}월)", header[1]).group(1)

    def col(label):
        for i, h in enumerate(header):
            if h == f"{month}_계_{label}":
                return i
        raise KeyError(label)

    c_total = col("총인구수")
    child_cols = [col(x) for x in ("0~4세", "5~9세", "10~14세")]
    elder_cols = [col(x) for x in ("65~69세", "70~74세", "75~79세", "80~84세", "85~89세", "90~94세", "95~99세", "100세 이상")]
    out = {}
    for r in rdr:
        if not r or not r[0]:
            continue
        m = re.match(r"(.+?)\s*\((\d{10})\)", r[0])
        if not m:
            continue
        name, code = m.group(1).strip(), m.group(2)
        if not (code.startswith("41") and code[2:4] != "00" and code[4:] == "000000"):
            continue
        sg = match_sigun(name)
        if not sg:
            continue
        n = lambda i: int(r[i].replace(",", ""))
        out[sg] = {
            "code": code,
            "total": n(c_total),
            "child_0_14": sum(n(i) for i in child_cols),
            "elder_65p": sum(n(i) for i in elder_cols),
        }
    return out, month


# ---------------------------------------------------------------- 경계
def decode_topo(topo):
    sx, sy = topo["transform"]["scale"]
    tx, ty = topo["transform"]["translate"]
    arcs = []
    for arc in topo["arcs"]:
        x = y = 0
        pts = []
        for dx, dy in arc:
            x += dx
            y += dy
            pts.append((x * sx + tx, y * sy + ty))
        arcs.append(pts)

    def ring(idxs):
        pts = []
        for k in idxs:
            a = arcs[~k][::-1] if k < 0 else arcs[k]
            if pts and pts[-1] == a[0]:
                pts.extend(a[1:])
            else:
                pts.extend(a)
        return pts

    return ring


def ring_area_km2(pts):
    # 소면적 등적 근사: 경도에 cos(위도) 보정 후 신발끈 공식
    lat0 = sum(p[1] for p in pts) / len(pts)
    k = 111.32
    xs = [(p[0] * math.cos(math.radians(lat0)) * k, p[1] * k) for p in pts]
    s = 0.0
    for i in range(len(xs)):
        x1, y1 = xs[i]
        x2, y2 = xs[(i + 1) % len(xs)]
        s += x1 * y2 - x2 * y1
    return abs(s) / 2


def load_boundaries(path):
    topo = json.load(open(path, encoding="utf-8"))
    ring = decode_topo(topo)
    obj = next(iter(topo["objects"].values()))
    features = []
    area = defaultdict(float)
    for g in obj["geometries"]:
        p = g["properties"]
        if not str(p.get("code", "")).startswith("31"):  # 통계청 시도코드 31 = 경기
            continue
        sg = match_sigun(p["name"])
        if not sg:
            raise SystemExit(f"경계 시군 매칭 실패: {p}")
        polys = [g["arcs"]] if g["type"] == "Polygon" else g["arcs"]
        coords = []
        a = 0.0
        for poly in polys:
            rings = [ring(r) for r in poly]
            a += ring_area_km2(rings[0]) - sum(ring_area_km2(r) for r in rings[1:])
            coords.append([[[round(x, 5), round(y, 5)] for x, y in r] for r in rings])
        area[sg] += a
        features.append({
            "type": "Feature",
            "properties": {"sigun": sg, "part": p["name"], "kostat_code": p["code"]},
            "geometry": {"type": "MultiPolygon", "coordinates": coords},
        })
    return {"type": "FeatureCollection", "features": features}, dict(area)


# ---------------------------------------------------------------- 버스정류장(선택)
def load_bus_stops(path):
    if not os.path.exists(path):
        return None
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "cp949"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    rdr = csv.reader(text.splitlines())
    header = next(rdr)
    hnorm = [norm(h) for h in header]
    c_city = next((i for i, h in enumerate(hnorm) if "도시명" in h), None)
    c_lat = next((i for i, h in enumerate(hnorm) if "위도" in h), None)
    c_lon = next((i for i, h in enumerate(hnorm) if "경도" in h), None)
    if c_city is None:
        return None
    cnt = Counter()
    for r in rdr:
        if len(r) <= c_city:
            continue
        city = r[c_city]
        sg = match_sigun(city)
        if sg and ("경기" in city or city.strip() in SIGUN or city.strip() in ALIAS):
            cnt[sg] += 1
    return dict(cnt), header


# ---------------------------------------------------------------- main
def main():
    os.makedirs(OUT, exist_ok=True)
    fac_rows, simple, unmatched = load_facilities(os.path.join(RAW, "mcst_public_sports_2024.xlsx"))
    pop, pop_month = load_population(os.path.join(RAW, "jumin_gg_age5_202609.csv"))
    geo, area = load_boundaries(os.path.join(RAW, "kostat_2018_municipalities_topo_simple.json"))
    bus = load_bus_stops(os.path.join(RAW, "bus_stops_gg_20251031.csv"))

    by_sigun_type = defaultdict(Counter)
    for r in fac_rows:
        by_sigun_type[r["sigun"]][r["type"]] += 1

    regions = []
    for sg in SIGUN:
        types = dict(sorted(by_sigun_type[sg].items(), key=lambda kv: -kv[1]))
        total = sum(types.values())
        p = pop[sg]
        rec = {
            "name": sg,
            "code": p["code"],
            "pop_total": p["total"],
            "pop_child_0_14": p["child_0_14"],
            "pop_elder_65p": p["elder_65p"],
            "area_km2": round(area[sg], 1),
            "facilities_total": total,
            "facilities_by_type": types,
            "facilities_life": sum(n for t, n in types.items() if t in LIFE_TYPES),
            "facilities_elder": sum(n for t, n in types.items() if t in ELDER_TYPES),
            "facilities_child": sum(n for t, n in types.items() if t in CHILD_TYPES),
            "simple_grounds": simple.get(sg),
        }
        if bus:
            rec["bus_stops"] = bus[0].get(sg, 0)
        regions.append(rec)

    missing_pop = [s for s in SIGUN if s not in pop]
    missing_area = [s for s in SIGUN if s not in area]
    if missing_pop or missing_area:
        raise SystemExit(f"누락: 인구 {missing_pop} / 면적 {missing_area}")

    meta = {
        "built_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "region_level": "경기도 시군(31개). 수원·성남·안양·안산·고양·용인의 일반구는 시 단위로 합산",
        "sources": {
            "facilities": {
                "name": "문화체육관광부 「2025년 전국 공공체육시설 현황(2024.12.31. 기준)」",
                "url": "https://www.mcst.go.kr/site/s_policy/dept/deptView.jsp?pSeq=2107&pDataCD=0417000000&pType=07",
                "reference_date": "2024-12-31",
                "license": "공공누리 제2유형(출처표시·상업적 이용금지)",
                "note": "23개 분류 시설별 시트의 '경기' 블록 행을 시군구·주소 기준으로 집계. 간이운동장(마을공터·도시공원 등)은 별도 표기하고 주요시설 수에는 넣지 않음.",
                "rows_counted": len(fac_rows),
                "rows_unmatched": len(unmatched),
            },
            "population": {
                "name": "행정안전부 주민등록 인구통계 — 연령별 인구현황(5세 구간)",
                "url": "https://jumin.mois.go.kr/ageStatMonth.do",
                "reference_date": pop_month,
                "license": "공공누리 제1유형(출처표시)",
                "note": "거주자+거주불명자+재외국민 합계(전체). 0~14세=아동, 65세 이상=고령자.",
            },
            "boundary": {
                "name": "통계청 센서스 행정구역경계 2018 (southkorea-maps 단순화본)",
                "url": "https://github.com/southkorea/southkorea-maps",
                "reference_date": "2018",
                "license": "원자료 통계청(SGIS) 공공데이터; 단순화·변환본은 southkorea-maps 저장소 배포",
                "note": "면적은 단순화 경계에서 계산한 근사치(㎢). 공식 행정구역 면적과 수 % 차이가 날 수 있음.",
            },
        },
        "indicators": {
            "facilities_per_10k": "주요 공공체육시설 수 ÷ 주민등록 총인구 × 10,000",
            "est_nearest_km": "0.5 × √(면적㎢ ÷ 시설 수). 시설이 면적에 고르게 퍼져 있다고 가정했을 때의 평균 최근접 직선거리(추정). 실제 경로·이동시간 아님",
            "vulnerable_ratio": "65세 이상(또는 0~14세) 인구 ÷ 총인구",
            "transit": "버스정류장 수 ÷ 면적㎢ (bus_stops 필드가 있을 때만)" if bus else "평가불가 — 경기도 대중교통 데이터 미확보",
        },
    }
    if bus:
        meta["sources"]["bus"] = {
            "name": "국토교통부 전국 버스정류장 위치정보 (2025-10-31)",
            "url": "https://www.data.go.kr/data/15067528/fileData.do",
            "reference_date": "2025-10-31",
            "license": "이용허락범위 제한 없음",
            "note": "도시명 컬럼으로 경기도 시군 매칭. BIS 미구축·TAGO 미연계 지자체는 0에 가깝게 나올 수 있음.",
        }

    json.dump({"meta": meta, "regions": regions}, open(os.path.join(OUT, "regions.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    json.dump(geo, open(os.path.join(OUT, "gg-sigun.geojson"), "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    with open(os.path.join(OUT, "facilities_gg.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["sigun", "type", "name", "address", "sheet"])
        w.writeheader()
        w.writerows(sorted(fac_rows, key=lambda r: (SIGUN.index(r["sigun"]), r["type"], r["name"])))

    print(f"시설 행 {len(fac_rows)}건 집계, 매칭 실패 {len(unmatched)}건")
    for u in unmatched[:20]:
        print("  미매칭:", u)
    print(f"인구 기준월 {pop_month}, 시군 {len(pop)}개 / 경계 폴리곤 {len(geo['features'])}개 / 면적 합계 {sum(area.values()):.0f}㎢")
    if bus:
        print(f"버스정류장 경기 합계 {sum(bus[0].values())}개 (헤더: {bus[1][:8]})")
    else:
        print("버스정류장 파일 없음 → transit 지표 평가불가")


if __name__ == "__main__":
    main()
